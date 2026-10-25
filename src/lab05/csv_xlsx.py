
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.lab05.json_csv import check_extension


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Преобразует CSV-файл в Excel.

    Первая строка CSV — заголовок.
    Лист называется Sheet1.
    Ширина столбцов зависит от содержимого, минимум 8 символов.
    Значения сохраняются как текст.
    """
    # Проверяем расширения файлов.
    check_extension(csv_path, ".csv")
    check_extension(xlsx_path, ".xlsx")

    # Читаем CSV.
    rows = []

    with open(csv_path, encoding="utf-8", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            if row:
                rows.append(row)

    # Проверяем, что есть заголовок и хотя бы одна строка данных.
    if len(rows) < 2:
        raise ValueError(
            "CSV пустой или в нём нет данных (только заголовок)"
        )

    # Находим максимальную длину текста в каждом столбце.
    widths = {}

    for row in rows:
        for index, cell in enumerate(row):
            if len(cell) > widths.get(index, 0):
                widths[index] = len(cell)

    # Создаём Excel-книгу и лист.
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    # Переносим строки из CSV в Excel.
    for row in rows:
        sheet.append(row)

    # Устанавливаем ширину столбцов.
    for index, width in widths.items():
        letter = get_column_letter(index + 1)
        sheet.column_dimensions[letter].width = max(8, width)

    # Создаём папку для результата, если её нет.
    output_path = Path(xlsx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Сохраняем Excel-файл.
    workbook.save(output_path)
    workbook.close()


if __name__ == "__main__":
    csv_to_xlsx(
        "data/lab05/samples/people.csv",
        "data/lab05/out/people.xlsx"
    )

    csv_to_xlsx(
        "data/lab05/samples/cities.csv",
        "data/lab05/out/cities.xlsx"
    )

    print("XLSX-файлы созданы.")
